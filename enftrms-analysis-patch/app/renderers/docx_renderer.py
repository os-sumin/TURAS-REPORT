from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.schemas.reports import GenerateReportRequest, RenderReportRequest, RenderedSection
from app.services.reports.tech_fee_prediction import predict_result_for_request
from app.services.reports.tech_fee_prediction.models import PredictionResult

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

SECTION_TITLES = {
    "SUMMARY": "Executive Summary",
    "MARKET": "Market Analysis",
    "COMPETITORS": "Competitor Analysis",
    "TECH_FEE": "Technology Fee",
    "RECOMMENDATIONS": "Recommendations",
}


class DocxRenderer:
    def render_generate_request(self, request: GenerateReportRequest, section_summaries: dict[str, str]) -> bytes:
        selected = request.report_context.sections or ["SUMMARY"]
        sections = [
            RenderedSection(
                code=code,
                title=SECTION_TITLES.get(code, code.replace("_", " ").title()),
                body=section_summaries.get(code, "No content generated."),
            )
            for code in selected
        ]

        tech_fee_prediction: PredictionResult | None = None
        if "TECH_FEE" in selected:
            try:
                tech_fee_prediction = predict_result_for_request(request)
            except Exception:
                tech_fee_prediction = None

        render_request = RenderReportRequest(
            report_id=request.report_id,
            output_format=request.output_format,
            render_context={
                "reportYear": request.report_context.report_year,
                "organizationName": request.organization.organization_name,
                "projectName": request.project.project_name,
                "contractName": request.contract.contract_name,
                "sections": [section.model_dump(by_alias=True) for section in sections],
                "financialInputs": request.financial_inputs,
                "metadata": {
                    "projectCode": request.project.project_code,
                    "contractNumber": request.contract.contract_number,
                    "licenseType": request.contract.license_type,
                    "agencyName": request.project.agency_name,
                },
            },
        )
        return self._render(render_request, tech_fee_prediction=tech_fee_prediction)

    def render_render_request(self, request: RenderReportRequest) -> bytes:
        return self._render(request)

    # ---------- core ----------

    def _render(self, request: RenderReportRequest, tech_fee_prediction: PredictionResult | None = None) -> bytes:
        document = Document()
        self._configure_document(document)
        self._add_cover(document, request)
        self._add_overview(document, request)
        self._add_sections(document, request.render_context.sections, tech_fee_prediction=tech_fee_prediction)
        self._add_financial_inputs(document, request.render_context.financial_inputs)
        self._add_metadata(document, request.render_context.metadata)

        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def _configure_document(self, document: Document) -> None:
        section = document.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

        styles = document.styles
        styles["Normal"].font.name = "Arial"
        styles["Normal"].font.size = Pt(10)
        styles["Heading 1"].font.name = "Arial"
        styles["Heading 1"].font.size = Pt(16)
        styles["Heading 2"].font.name = "Arial"
        styles["Heading 2"].font.size = Pt(12)

    def _add_cover(self, document: Document, request: RenderReportRequest) -> None:
        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.add_run("EnFTRMS AI Commercialization Report")
        title_run.bold = True
        title_run.font.size = Pt(22)

        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.add_run(request.report_id).font.size = Pt(10)

        organization = request.render_context.organization_name or "-"
        project = request.render_context.project_name or "-"
        contract = request.render_context.contract_name or "-"
        year = request.render_context.report_year or "-"

        summary = document.add_paragraph()
        summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
        summary.add_run(f"{year} | {organization} | {project} | {contract}").italic = True

        document.add_paragraph()

    def _add_overview(self, document: Document, request: RenderReportRequest) -> None:
        document.add_heading("Overview", level=1)
        rows = [
            ("Report ID", request.report_id),
            ("Organization", request.render_context.organization_name or "-"),
            ("Project", request.render_context.project_name or "-"),
            ("Contract", request.render_context.contract_name or "-"),
            ("Report Year", str(request.render_context.report_year or "-")),
            ("Sections", ", ".join(section.code for section in request.render_context.sections) or "-"),
        ]

        table = document.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.style = "Table Grid"
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = value

        document.add_paragraph()

    def _add_sections(
        self,
        document: Document,
        sections: list[RenderedSection],
        tech_fee_prediction: PredictionResult | None = None,
    ) -> None:
        for section in sections:
            document.add_heading(section.title or section.code, level=1)
            for paragraph in self._split_paragraphs(section.body):
                document.add_paragraph(paragraph)
            if section.code == "TECH_FEE" and tech_fee_prediction is not None:
                self._add_tech_fee_table(document, tech_fee_prediction)

    def _add_tech_fee_table(self, document: Document, result: PredictionResult) -> None:
        # 1) 차원별 점수 표
        document.add_heading("5대 차원별 점수", level=2)
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["코드", "차원", "점수", "만점"]
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h
        for d in result.dimensions:
            cells = table.add_row().cells
            cells[0].text = d.code
            cells[1].text = d.label
            cells[2].text = f"{d.score:.1f}"
            cells[3].text = f"{d.max_score:.0f}"

        # 2) 총점·등급 강조
        document.add_paragraph()
        emphasis = document.add_paragraph()
        run = emphasis.add_run(
            f"총점 {result.total_score:.1f} / 100점 — 등급: {result.grade.label} ({result.grade.code})"
        )
        run.bold = True
        run.font.size = Pt(11)
        document.add_paragraph(f"추천 조치: {result.grade.action}")
        document.add_paragraph(f"룰 버전: {result.rule_version}")
        document.add_paragraph()

    def _add_financial_inputs(self, document: Document, rows: list[dict[str, Any]]) -> None:
        if not rows:
            return

        document.add_heading("Financial Inputs", level=1)
        first_row = rows[0]
        if not first_row:
            return

        headers = [str(key) for key in first_row.keys()]
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header

        for row in rows:
            cells = table.add_row().cells
            for index, header in enumerate(headers):
                cells[index].text = self._stringify(row.get(header, ""))

        document.add_paragraph()

    def _add_metadata(self, document: Document, metadata: dict[str, Any]) -> None:
        if not metadata:
            return

        document.add_heading("Metadata", level=1)
        for key, value in metadata.items():
            document.add_paragraph(f"{key}: {self._stringify(value)}")

    def _split_paragraphs(self, body: str) -> list[str]:
        chunks = [chunk.strip() for chunk in body.split("\n") if chunk.strip()]
        return chunks or ["-"]

    def _stringify(self, value: Any) -> str:
        if value is None:
            return "-"
        if isinstance(value, list):
            return ", ".join(self._stringify(item) for item in value) or "-"
        if isinstance(value, dict):
            return ", ".join(f"{key}={self._stringify(val)}" for key, val in value.items()) or "-"
        return str(value)
