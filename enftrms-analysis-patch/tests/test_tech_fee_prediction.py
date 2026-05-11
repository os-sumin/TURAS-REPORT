from datetime import datetime

import pytest

from app.schemas.reports import (
    ContractContext,
    GenerateReportRequest,
    ReportContext,
)
from app.services.reports.heuristic_sections import build_heuristic_sections
from app.services.reports.tech_fee_prediction import (
    PredictionContext,
    TfeeHistory,
    run,
)
from app.services.reports.tech_fee_prediction.dimensions import d3_capital_efficiency
from app.services.reports.tech_fee_prediction.evaluator import evaluate
from app.services.reports.tech_fee_prediction.input_adapter import build_context
from app.services.reports.tech_fee_prediction.rules import (
    Bucket,
    Rule,
    Threshold,
    load_cached,
)
from app.services.reports.tech_fee_prediction.section_writer import (
    predict_and_write_section,
)


# ----- evaluator -----

def test_evaluator_threshold():
    rule = Rule(name="t", max=5, type="threshold", thresholds=[
        Threshold(gte=0.2, score=6),
        Threshold(gte=0.0, score=1),
    ])
    assert evaluate(rule, 0.25) == 5      # capped at max
    assert evaluate(rule, 0.0) == 1


def test_evaluator_bucket_inclusive_lo_exclusive_hi():
    rule = Rule(name="b", max=5, type="bucket", buckets=[
        Bucket(range=[12, 24], score=5),
    ])
    assert evaluate(rule, 12) == 5
    assert evaluate(rule, 24) == 0


def test_evaluator_binary_str_truthy():
    rule = Rule(name="bn", max=6, type="binary", score_when_true=6)
    assert evaluate(rule, "Y") == 6
    assert evaluate(rule, "n") == 0


# ----- D3 (full impl) -----

def test_d3_high_score():
    dim = load_cached().dimensions["D3_capital_efficiency"]
    ctx = PredictionContext(
        sbjt_id="S", orgn_id="O", pred_at=datetime.now(),
        tfee_history=TfeeHistory(
            has_past_tfee=True, recovery_rate_percentile=0.95,
            project_recovery_rate=0.25, months_to_first_tfee=8, consecutive_years=3,
        ),
    )
    r = d3_capital_efficiency.calculate(ctx, dim)
    assert 19.0 <= r.score <= 20.0


def test_d3_low_score():
    dim = load_cached().dimensions["D3_capital_efficiency"]
    ctx = PredictionContext(
        sbjt_id="S", orgn_id="O", pred_at=datetime.now(),
        tfee_history=TfeeHistory(),
    )
    r = d3_capital_efficiency.calculate(ctx, dim)
    assert 0.0 <= r.score <= 1.5     # project_recovery_rate=0 → gte 0 룰 매칭으로 1점


# ----- engine -----

def test_engine_produces_five_dimensions():
    ctx = PredictionContext(
        sbjt_id="S001", orgn_id="O001", pred_at=datetime.now(),
        tfee_history=TfeeHistory(has_past_tfee=True, recovery_rate_percentile=0.5,
                                  project_recovery_rate=0.15, months_to_first_tfee=18,
                                  consecutive_years=2),
    )
    r = run(ctx)
    assert len(r.dimensions) == 5
    assert r.grade.code in {"VERY_HIGH", "HIGH", "MID", "LOW", "VERY_LOW"}
    assert 0.0 <= r.total_score <= 100.0


def test_engine_empty_context_is_very_low():
    ctx = PredictionContext(sbjt_id="S", orgn_id="O", pred_at=datetime.now())
    r = run(ctx)
    assert r.grade.code == "VERY_LOW"


# ----- input adapter -----

def _request(analysis: dict) -> GenerateReportRequest:
    return GenerateReportRequest(
        report_id="AIR-T1",
        report_context=ReportContext(sections=["TECH_FEE"]),
        contract=ContractContext(contract_name="Test Contract"),
        organization={"organizationId": 101, "organizationName": "테스트기업"},
        project={"projectId": 1001, "projectCode": "PJT-T", "projectName": "테스트과제"},
        analysis=analysis,
    )


def test_adapter_accepts_tfee_history_block():
    req = _request({"techFee": {
        "tfeeHistory": {"hasPastTfee": True, "projectRecoveryRate": 0.15,
                        "recoveryRatePercentile": 0.6, "monthsToFirstTfee": 15,
                        "consecutiveYears": 2},
    }})
    ctx = build_context(req)
    assert ctx is not None
    assert ctx.tfee_history.has_past_tfee is True
    assert ctx.tfee_history.project_recovery_rate == 0.15
    assert ctx.company_name == "테스트기업"


def test_adapter_accepts_excel_row_form():
    req = _request({"techFee": {"row": {
        "ttlPayGvstmAm": 100_000_000, "ttlUseGvstmAm": 99_000_000,
        "salesOccurYn": "Y", "tfeeAm": 200_000, "techCtrbPt": 0.5,
        "baseYear": 2026,
    }}})
    ctx = build_context(req)
    assert ctx is not None
    assert ctx.tfee_history.has_past_tfee is True
    assert ctx.tfee_history.project_recovery_rate == pytest.approx(200_000 / 99_000_000)


def test_adapter_derives_financials_from_financial_inputs():
    req = GenerateReportRequest(
        report_id="AIR-T2",
        report_context=ReportContext(sections=["TECH_FEE"]),
        contract=ContractContext(contract_name="C"),
        organization={"organizationId": 1, "organizationName": "X"},
        project={"projectId": 1, "projectCode": "P", "projectName": "Y"},
        analysis={"techFee": {}},
        financial_inputs=[
            {"year": 2024, "revenue": 1000, "operatingIncome": 100},
            {"year": 2025, "revenue": 1200, "operatingIncome": 180},
        ],
    )
    ctx = build_context(req)
    assert ctx.financials.sales_growth_rate == pytest.approx(0.2)
    assert ctx.financials.operating_margin == pytest.approx(0.15)


# ----- end-to-end through heuristic_sections -----

def test_heuristic_tech_fee_uses_prediction():
    req = _request({"techFee": {
        "tfeeHistory": {"hasPastTfee": True, "projectRecoveryRate": 0.20,
                        "recoveryRatePercentile": 0.8, "monthsToFirstTfee": 10,
                        "consecutiveYears": 2},
    }})
    result = build_heuristic_sections(req, ["TECH_FEE"])
    body = result["TECH_FEE"]
    assert "기술료 납부 가능성" in body
    assert "5대 차원별 점수" in body
    assert "D3" in body
    # 기존 fallback 문구가 절대 나오면 안됨
    assert "Spring payload에 기술료 분석 데이터" not in body


def test_heuristic_tech_fee_falls_back_when_no_payload():
    req = _request({})
    # techFee/financial_inputs 모두 없음 → input_adapter 가 None 반환 →
    # _stringify_block 의 default_text 가 노출
    result = build_heuristic_sections(req, ["TECH_FEE"])
    assert "포함되지 않았습니다" in result["TECH_FEE"]


def test_predict_and_write_section_direct():
    req = _request({"techFee": {
        "tfeeHistory": {"hasPastTfee": True, "projectRecoveryRate": 0.15,
                        "recoveryRatePercentile": 0.6, "monthsToFirstTfee": 15,
                        "consecutiveYears": 2},
        "externalMetrics": {"g2bBidCount": 5},
        "newsEventCounts": {"PRODUCT_LAUNCH": 1, "CERTIFICATION": 1},
    }})
    body = predict_and_write_section(req)
    assert body is not None
    assert "총점" in body
    assert "/100점" in body


# ----- news event classifier -----

def test_classifier_detects_product_launch_and_certification():
    from app.services.reports.tech_fee_prediction.news_event_classifier import classify_articles
    counts = classify_articles([
        {"title": "넥스트바이오, AI 진단 플랫폼 상용화 출시"},
        {"title": "식약처 혁신의료기기 지정 — 인증 획득"},
        {"title": "서울대병원과 공급계약 체결"},
    ])
    assert counts.get("PRODUCT_LAUNCH", 0) >= 1
    assert counts.get("CERTIFICATION", 0) >= 1
    assert counts.get("SUPPLY_CONTRACT", 0) >= 1


def test_classifier_negative_priority():
    from app.services.reports.tech_fee_prediction.news_event_classifier import classify_articles
    # 같은 기사에 출시 + 소송 키워드가 같이 있으면 부정 우선
    counts = classify_articles([
        {"title": "신제품 출시 소식과 함께 진행되는 특허 소송 분쟁"},
    ])
    assert counts.get("LAWSUIT", 0) == 1
    assert counts.get("PRODUCT_LAUNCH", 0) == 0


def test_classifier_unmatched_article_does_not_count():
    from app.services.reports.tech_fee_prediction.news_event_classifier import classify_articles
    counts = classify_articles([
        {"title": "올해 분기별 매출 변동에 대한 일반 분석"},
    ])
    assert counts == {}


def test_adapter_classifies_articles_when_no_explicit_counts():
    req = _request({"techFee": {
        "tfeeHistory": {"hasPastTfee": True, "projectRecoveryRate": 0.1},
        "articles": [
            {"title": "공급계약 체결로 매출 확대"},
            {"title": "식약처 허가 획득"},
            {"title": "투자유치 시리즈B 완료"},
        ],
    }})
    ctx = build_context(req)
    assert ctx is not None
    assert ctx.news_event_counts.get("SUPPLY_CONTRACT", 0) >= 1
    assert ctx.news_event_counts.get("CERTIFICATION", 0) >= 1
    assert ctx.news_event_counts.get("INVESTMENT", 0) >= 1


def test_adapter_explicit_counts_override_articles():
    req = _request({"techFee": {
        "newsEventCounts": {"PRODUCT_LAUNCH": 3},
        "articles": [{"title": "공급계약 체결"}],  # 무시되어야 함
    }})
    ctx = build_context(req)
    assert ctx.news_event_counts == {"PRODUCT_LAUNCH": 3}


# ----- DOCX renderer table -----

def test_docx_renderer_embeds_tech_fee_table():
    from app.services.reports.service import ReportService
    import base64
    from io import BytesIO
    from docx import Document

    req = _request({"techFee": {
        "tfeeHistory": {"hasPastTfee": True, "projectRecoveryRate": 0.18,
                        "recoveryRatePercentile": 0.72, "monthsToFirstTfee": 14,
                        "consecutiveYears": 2},
        "financials": {"salesGrowthRate": 0.18, "operatingMargin": 0.06},
        "externalMetrics": {"g2bBidCount": 5},
        "articles": [
            {"title": "공급계약 체결"},
            {"title": "인증 획득"},
        ],
    }})
    # TECH_FEE 섹션이 sections 에 포함되어 있어야 표가 그려짐
    req.report_context.sections = ["SUMMARY", "TECH_FEE", "RECOMMENDATIONS"]

    resp = ReportService().generate(req)
    assert resp.status == "COMPLETED"

    doc = Document(BytesIO(base64.b64decode(resp.content_base64)))
    table_texts = [
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    ]
    # 5대 차원 코드가 표에 모두 나타나야 함
    assert any("D1" == t for t in table_texts)
    assert any("D2" == t for t in table_texts)
    assert any("D3" == t for t in table_texts)
    assert any("D4" == t for t in table_texts)
    assert any("D5" == t for t in table_texts)

    # 등급/추천조치 단락 존재
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "총점" in all_text
    assert "추천 조치" in all_text
